#!/usr/bin/env python3
"""Generate design-doc Word documents from markdown source.

Uses design-doc-base.docx (or chat docx fallback) as a style template.
Supports headings, paragraphs, bullet lists, fenced code blocks, and pipe tables.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:
    import subprocess

    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx', '-q'])
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
DEFAULT_TEMPLATE = ROOT / 'docs' / 'design_docs' / 'templates' / 'design-doc-base.docx'
FALLBACK_TEMPLATE = ROOT / 'docs' / 'chat-operation-scope-capabilities-2026-09-03.docx'


def clear_document_body(doc: Document) -> None:
    body = doc.element.body
    sect_pr = None
    for child in list(body):
        if child.tag == qn('w:sectPr'):
            sect_pr = child
            continue
        if child.tag in (qn('w:p'), qn('w:tbl')):
            body.remove(child)
    if sect_pr is not None and sect_pr.getparent() is None:
        body.append(sect_pr)


def add_runs_with_bold(paragraph: Paragraph, text: str) -> None:
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def add_paragraph(doc: Document, text: str, style: str = 'Normal') -> Paragraph:
    paragraph = doc.add_paragraph(style=style)
    add_runs_with_bold(paragraph, text)
    return paragraph


def add_bullet(doc: Document, text: str, level: int = 0) -> Paragraph:
    style = 'List Bullet' if level == 0 else 'List Bullet 2'
    try:
        paragraph = doc.add_paragraph(style=style)
    except KeyError:
        paragraph = doc.add_paragraph(style='List Bullet')
    add_runs_with_bold(paragraph, text)
    return paragraph


def parse_table_row(line: str) -> list[str] | None:
    stripped = line.strip()
    if not stripped.startswith('|'):
        return None
    cells = [cell.strip() for cell in stripped.strip('|').split('|')]
    return cells


def is_separator_row(cells: list[str]) -> bool:
    return all(re.fullmatch(r':?-{3,}:?', cell.replace(' ', '')) for cell in cells if cell != '')


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> Table:
    col_count = max(len(headers), max((len(r) for r in rows), default=0))
    if col_count < 1:
        col_count = 1

    def pad(row: list[str]) -> list[str]:
        padded = row[:col_count]
        while len(padded) < col_count:
            padded.append('')
        return padded

    table_rows = [pad(headers)] + [pad(r) for r in rows]
    table = doc.add_table(rows=len(table_rows), cols=col_count)
    table.style = 'Table Grid'

    for ri, row in enumerate(table_rows):
        for ci, value in enumerate(row):
            cell = table.rows[ri].cells[ci]
            cell.text = value
            if ri == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
    doc.add_paragraph('')
    return table


def parse_markdown_blocks(lines: list[str]) -> list[dict]:
    blocks: list[dict] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == '---':
            blocks.append({'type': 'hr'})
            i += 1
            continue

        if stripped.startswith('```'):
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            blocks.append({'type': 'code', 'text': '\n'.join(code_lines).rstrip()})
            continue

        if stripped.startswith('#'):
            level = len(stripped) - len(stripped.lstrip('#'))
            text = stripped[level:].strip()
            blocks.append({'type': 'heading', 'level': min(level, 3), 'text': text})
            i += 1
            continue

        if stripped.startswith('|'):
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            parsed = [parse_table_row(row) for row in table_lines]
            parsed = [row for row in parsed if row is not None]
            if parsed:
                headers = parsed[0]
                data_rows = parsed[1:]
                if data_rows and is_separator_row(data_rows[0]):
                    data_rows = data_rows[1:]
                blocks.append({'type': 'table', 'headers': headers, 'rows': data_rows})
            continue

        if stripped.startswith('- ') or stripped.startswith('* '):
            items: list[str] = []
            while i < len(lines):
                item_line = lines[i].strip()
                if item_line.startswith('- ') or item_line.startswith('* '):
                    items.append(item_line[2:].strip())
                    i += 1
                elif item_line.startswith('  - ') or item_line.startswith('  * '):
                    items.append('  ' + item_line.strip()[2:].strip())
                    i += 1
                else:
                    break
            blocks.append({'type': 'list', 'items': items})
            continue

        para_lines: list[str] = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith('#') or nxt.startswith('|') or nxt.startswith('```') or nxt == '---' or nxt.startswith('- ') or nxt.startswith('* '):
                break
            para_lines.append(nxt)
            i += 1
        blocks.append({'type': 'paragraph', 'text': ' '.join(para_lines)})

    return blocks


def heading_style(level: int) -> str:
    return {1: 'Heading 1', 2: 'Heading 2', 3: 'Heading 3'}.get(level, 'Heading 3')


def render_blocks(doc: Document, blocks: list[dict]) -> None:
    for block in blocks:
        block_type = block['type']
        if block_type == 'heading':
            add_paragraph(doc, block['text'], heading_style(block['level']))
        elif block_type == 'paragraph':
            add_paragraph(doc, block['text'])
        elif block_type == 'hr':
            doc.add_paragraph('')
        elif block_type == 'code':
            for code_line in block['text'].split('\n') or ['']:
                paragraph = doc.add_paragraph(style='Normal')
                run = paragraph.add_run(code_line)
                run.font.name = 'Consolas'
        elif block_type == 'list':
            for item in block['items']:
                if item.startswith('  '):
                    add_bullet(doc, item.strip(), level=1)
                else:
                    add_bullet(doc, item)
        elif block_type == 'table':
            add_table(doc, block['headers'], block['rows'])


def ensure_template(template_path: Path) -> Path:
    if template_path.exists():
        return template_path

    template_path.parent.mkdir(parents=True, exist_ok=True)
    if not FALLBACK_TEMPLATE.exists():
        raise FileNotFoundError(
            f'Missing template {template_path} and fallback {FALLBACK_TEMPLATE}'
        )

    fallback = Document(str(FALLBACK_TEMPLATE))
    clear_document_body(fallback)
    add_paragraph(fallback, 'Design Document Template', 'Heading 1')
    add_paragraph(fallback, 'Template shell for design_docs Word output.')
    fallback.save(str(template_path))
    return template_path


def build_document(input_path: Path, output_path: Path, template_path: Path) -> Path:
    if not input_path.exists():
        raise FileNotFoundError(f'Missing markdown input: {input_path}')

    template = ensure_template(template_path)
    doc = Document(str(template))
    clear_document_body(doc)

    text = input_path.read_text(encoding='utf-8')
    blocks = parse_markdown_blocks(text.splitlines())
    render_blocks(doc, blocks)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def main() -> int:
    parser = argparse.ArgumentParser(description='Generate design doc DOCX from markdown.')
    parser.add_argument('input', type=Path, help='Input markdown file')
    parser.add_argument('output', type=Path, nargs='?', default=None, help='Output docx file')
    parser.add_argument(
        '--template',
        type=Path,
        default=DEFAULT_TEMPLATE,
        help='DOCX style template',
    )
    args = parser.parse_args()

    input_path = args.input if args.input.is_absolute() else ROOT / args.input
    output_path = args.output
    if output_path is None:
        output_path = input_path.with_suffix('.docx')
    elif not output_path.is_absolute():
        output_path = ROOT / output_path

    template_path = args.template if args.template.is_absolute() else ROOT / args.template

    result = build_document(input_path, output_path, template_path)
    print(f'Generated: {result}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
