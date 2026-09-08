#!/usr/bin/env python3
"""Apply tracked edits to attendance operation-scope docx.

Deletions: strikethrough. Additions: red text.
"""

from __future__ import annotations

import copy
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:
    import subprocess

    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx', '-q'])
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import RGBColor
    from docx.table import Table
    from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
DOCX_PATH = ROOT / 'docs' / 'attendance-operation-scope-capabilities-2026-09-04.docx'
RED = RGBColor(0xFF, 0, 0)


def get_tables(doc: Document) -> list[Table]:
    return [Table(child, doc) for child in doc.element.body.iterchildren() if child.tag.endswith('tbl')]


def find_table_containing(doc: Document, needle: str) -> Table | None:
    for table in get_tables(doc):
        for row in table.rows:
            for cell in row.cells:
                if needle in cell.text:
                    return table
    return None


def clear_cell(cell) -> None:
    cell.text = ''


def add_strike(paragraph: Paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.font.strike = True


def add_red(paragraph: Paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.font.color.rgb = RED


def set_paragraph_tracked_full(paragraph: Paragraph, old_text: str, new_text: str) -> None:
    paragraph.clear()
    if old_text:
        add_strike(paragraph, old_text)
    if old_text and new_text:
        paragraph.add_run('\n')
    if new_text:
        add_red(paragraph, new_text)


def set_cell_tracked_full(cell, old_text: str, new_text: str) -> None:
    clear_cell(cell)
    set_paragraph_tracked_full(cell.paragraphs[0], old_text, new_text)


def append_cell_tracked(cell, strike_text: str = '', red_text: str = '', prefix: str = '\n') -> None:
    paragraph = cell.paragraphs[0]
    if strike_text:
        if paragraph.text:
            paragraph.add_run(prefix)
        add_strike(paragraph, strike_text)
    if red_text:
        paragraph.add_run(prefix)
        add_red(paragraph, red_text)


def replace_in_cell_tracked(cell, old_fragment: str, red_replacement: str) -> None:
    text = cell.text
    if old_fragment not in text:
        append_cell_tracked(cell, strike_text='', red_text=red_replacement)
        return
    before, after = text.split(old_fragment, 1)
    clear_cell(cell)
    paragraph = cell.paragraphs[0]
    if before:
        paragraph.add_run(before)
    add_strike(paragraph, old_fragment)
    add_red(paragraph, red_replacement)
    if after:
        paragraph.add_run(after)


def set_cell_simple(cell, text: str, red: bool = False) -> None:
    clear_cell(cell)
    paragraph = cell.paragraphs[0]
    if red:
        add_red(paragraph, text)
    else:
        paragraph.add_run(text)


def insert_paragraph_before(paragraph: Paragraph, text: str = '', style: str = 'Normal', red: bool = False) -> Paragraph:
    new_p = OxmlElement('w:p')
    paragraph._element.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = style
    if text:
        if red:
            add_red(new_para, text)
        else:
            new_para.add_run(text)
    return new_para


def insert_paragraph_after(paragraph: Paragraph, text: str = '', style: str = 'Normal', red: bool = False) -> Paragraph:
    new_p = OxmlElement('w:p')
    paragraph._element.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.style = style
    if text:
        if red:
            add_red(new_para, text)
        else:
            new_para.add_run(text)
    return new_para


def insert_table_after(paragraph: Paragraph, doc: Document, rows: int, cols: int) -> Table:
    table = doc.add_table(rows=rows, cols=cols)
    table.style = 'Table Grid'
    tbl_element = table._element
    tbl_element.getparent().remove(tbl_element)
    paragraph._element.addnext(tbl_element)
    return table


def delete_table_rows(table: Table, start_index: int) -> None:
    while len(table.rows) > start_index:
        table._tbl.remove(table.rows[start_index]._tr)


def add_table_row(table: Table, after_index: int | None = None) -> None:
    new_row = table.add_row()
    if after_index is None:
        return
    new_tr = table.rows[-1]._tr
    table._tbl.remove(new_tr)
    table.rows[after_index]._tr.addnext(new_tr)
    for cell in table.rows[after_index + 1].cells:
        cell.text = ''


def fill_table_header(table: Table, headers: list[str]) -> None:
    for ci, header in enumerate(headers):
        set_cell_simple(table.rows[0].cells[ci], header)
        for run in table.rows[0].cells[ci].paragraphs[0].runs:
            run.bold = True


def replace_admin_table(table: Table) -> None:
    old_rows = []
    for ri in range(1, len(table.rows)):
        old_rows.append((
            table.rows[ri].cells[0].text,
            table.rows[ri].cells[1].text,
        ))

    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)

    set_cell_simple(table.rows[0].cells[0], 'Topic')
    set_cell_simple(table.rows[0].cells[1], 'See')
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True

    combined_old = '\n\n'.join(f'{a}\n{b}' for a, b in old_rows)
    new_rows = [
        (
            'Bypass admin types (Family A)',
            'docs/design_docs/admin-access-types-reference-2026-09-05.md (and matching .docx). '
            'These users bypass operation/scope rows for SCHOOL_ATTENDANCES unless a separate business rule applies.',
        ),
        (
            'ADMIN in scope column (Family B)',
            'Access Profile ADMIN scope (SCP_ADMIN) only — incremental "+" rows in operation tables below. '
            'Not the same as bypass section/global admins.',
        ),
    ]

    for index, (topic, see_text) in enumerate(new_rows):
        row = table.add_row()
        old_a = combined_old if index == 0 else ''
        old_b = combined_old if index == 0 else ''
        set_cell_tracked_full(row.cells[0], old_a, topic)
        set_cell_tracked_full(row.cells[1], old_b, see_text)


def rebuild_delete_table(doc: Document, old_table: Table) -> None:
    parent = old_table._element.getparent()
    index = list(parent).index(old_table._element)

    old_snapshot = []
    for row in old_table.rows:
        old_snapshot.append(' | '.join(cell.text for cell in row.cells))

    parent.remove(old_table._element)

    new_table = doc.add_table(rows=3, cols=3)
    new_table.style = 'Table Grid'
    new_element = new_table._element
    new_element.getparent().remove(new_element)
    parent.insert(index, new_element)

    fill_table_header(new_table, ['Operation', 'Scope', 'What user can do and see'])
    old_text = '\n'.join(old_snapshot)
    rows_data = [
        ('DELETE', 'USER / OWNER', 'No Access.'),
        (
            'DELETE',
            'DEPARTMENT / DIVISION / ORGANIZATION / ADMIN',
            '+ Can delete files attached to attendance details (Matrix + Manage Session where files are exposed).',
        ),
    ]
    for ri, (op, scope, desc) in enumerate(rows_data, start=1):
        set_cell_simple(new_table.rows[ri].cells[0], op, red=True)
        set_cell_simple(new_table.rows[ri].cells[1], scope, red=True)
        set_cell_tracked_full(new_table.rows[ri].cells[2], old_text if ri == 1 else '', desc)


def split_scope_row(table: Table, row_index: int, org_desc: str, admin_desc: str, operation: str = '') -> None:
    row = table.rows[row_index]
    op = operation or row.cells[0].text.strip()
    old_scope = row.cells[1].text
    old_desc = row.cells[2].text if len(row.cells) > 2 else row.cells[1].text

    set_cell_tracked_full(row.cells[1], old_scope, 'OWNER / DEPARTMENT / DIVISION')
    if len(row.cells) > 2:
        append_cell_tracked(row.cells[2], red_text='Applies to non-bypass users only.')

    add_table_row(table, after_index=row_index)
    org_row = table.rows[row_index + 1]
    set_cell_simple(org_row.cells[0], op, red=True)
    set_cell_simple(org_row.cells[1], 'ORGANIZATION', red=True)
    set_cell_simple(org_row.cells[2], org_desc, red=True)

    add_table_row(table, after_index=row_index + 1)
    admin_row = table.rows[row_index + 2]
    set_cell_simple(admin_row.cells[0], op, red=True)
    set_cell_simple(admin_row.cells[1], 'ADMIN', red=True)
    set_cell_simple(admin_row.cells[2], admin_desc, red=True)


def split_export_print_table(table: Table) -> None:
    row = table.rows[2]
    old_scope = row.cells[1].text
    old_desc = row.cells[2].text
    set_cell_tracked_full(row.cells[1], old_scope, 'DEPARTMENT / DIVISION / ORGANIZATION')
    set_cell_simple(row.cells[2], old_desc)

    add_table_row(table, after_index=2)
    admin_row = table.rows[3]
    op = table.rows[1].cells[0].text
    set_cell_simple(admin_row.cells[0], op, red=True)
    set_cell_simple(admin_row.cells[1], 'ADMIN', red=True)
    set_cell_simple(admin_row.cells[2], f'No additional {op} extras beyond ORGANIZATION.', red=True)


def build_admin_overrides_table(doc: Document, anchor: Paragraph) -> Table:
    intro_old = anchor.text
    intro_new = (
        'Attendance section admin overrides apply to bypass admins (Family A — see admin-access-types-reference). '
        'They do not replace READ / READ_ALL / UPDATE rows; they bypass selected business locks. '
        'Implementation: schoolAdminAccessService.isAttendancesAdminViewerAsync, userCanOpenAttendanceMatrix, '
        'userCanMarkAttendanceExcused (excuses follow ORGANIZATION UPDATE scope, not A6 override).'
    )
    set_paragraph_tracked_full(anchor, intro_old, intro_new)

    table = insert_table_after(anchor, doc, rows=13, cols=4)
    fill_table_header(table, ['#', 'Feature / override', 'Applies to bypass admins?', 'In override form?'])
    rows = [
        ('A1', 'Edit attendance on locked session', 'Yes', 'Yes'),
        ('A2', 'Edit attendance after completed-session deadline', 'Yes', 'Yes'),
        ('A3', 'View attendance change history (when restricted below ORG READ_ALL)', 'Yes', 'Yes'),
        ('A4', 'Attendance section admin viewer tools on matrix', 'Yes', 'Yes'),
        ('A5', 'Open Attendance Matrix without normal READ/UPDATE', 'Yes', 'Yes'),
        ('A6', 'Mark excuse flags / excuse notes', 'N/A — scope-based at ORGANIZATION UPDATE', 'No'),
        ('A7', 'Bypass enrollment window', 'No', 'Not attendance admin override'),
        ('A8', 'Bypass makeup-required original session', 'No', 'Not attendance admin override'),
        ('A9', 'Bypass enrollment-office N/A lock', 'No (SCHOOL_CLASSES class admin)', 'Not attendance admin override'),
        ('A10', 'Bypass rolling enrollment capacity', 'No', 'Not attendance admin override'),
        ('A11', 'Lock / unlock session metadata', 'No (class/session admin)', 'Not attendance admin override'),
        ('A12', 'Full field visibility regardless of scope', 'Optional bypass', 'Note only — see admin reference'),
    ]
    for ri, row_data in enumerate(rows, start=1):
        for ci, value in enumerate(row_data):
            set_cell_simple(table.rows[ri].cells[ci], value, red=True)
    doc.add_paragraph('')
    return table


def apply_updates(doc: Document) -> None:
    tables = get_tables(doc)

    # --- Admin intro paragraphs after P15 ---
    admin_heading = next(p for p in doc.paragraphs if p.text.strip() == 'ADMINS And their privileges:')
    p1 = insert_paragraph_after(
        admin_heading,
        'See docs/design_docs/admin-access-types-reference-2026-09-05.md (and matching .docx) for all bypass admin types (Family A). ',
        red=True,
    )
    p2 = insert_paragraph_after(
        p1,
        'Operation/scope rows in this document apply to non-bypass users only. ADMIN in the scope column means Access Profile ADMIN scope (Family B / SCP_ADMIN) only — not bypass section or global admins. Attendance-specific bypass behaviors are listed under Admin and Other Scopes overrides below.',
        red=True,
    )
    insert_paragraph_after(
        p2,
        'Related documents: docs/design_docs/admin-access-types-reference-2026-09-05.md (bypass vs ADMIN scope; developer API for admin checks); docs/manage-session-access-architecture-report-2026-09-04.md (Manage Session runtime split).',
        red=True,
    )

    replace_admin_table(tables[2])

    # --- Scope ladder note before Access Profile Definitions ---
    access_heading = next(
        p for p in doc.paragraphs
        if p.text.strip() == 'Access Profile Definitions' and p.style.name == 'Heading 2'
    )
    insert_paragraph_before(
        access_heading,
        '+ = additional visibility or actions at this scope level. Effective access is cumulative unless stated otherwise. Operation tables apply to non-bypass users only.',
        red=True,
    )

    # --- Table 1: Pages route gate ---
    pages = tables[1]
    matrix_cell = pages.rows[1].cells[1]
    replace_in_cell_tracked(
        matrix_cell,
        'Route gate: SCHOOL_ATTENDANCES UPDATE.',
        'Route gate (target): SCHOOL_ATTENDANCES READ to open page; UPDATE for mutations. Current runtime: page and APIs still require UPDATE until promotion work is done.',
    )

    # --- Table 3: READ split ORG/ADMIN ---
    read_table = tables[3]
    split_scope_row(
        read_table,
        2,
        'Open Attendance Matrix page only; no attendance field visibility unless READ_ALL is also granted; no mutations unless UPDATE is granted.',
        'Open Attendance Matrix page only; no additional READ extras at this tier (ADMIN scope / Family B).',
        operation='READ',
    )
    note_row = read_table.rows[-1]
    append_cell_tracked(note_row.cells[2 if len(note_row.cells) > 2 else 1], red_text='Applies to non-bypass users only.')

    # --- Table 4: READ_ALL ADMIN row + notes ---
    read_all = tables[4]
    admin_row = read_all.rows[7]
    old_admin_desc = admin_row.cells[2].text
    new_admin_desc = (
        'On the "Attendance Matrix" page: + Can view attendance on locked sessions '
        '(read-only visibility where lower scopes cannot see locked session cells). '
        'Bypass-only items (edit locked session, admin viewer tools, open without READ/UPDATE, full field visibility) '
        'are listed in Admin overrides A1, A4, A5, A12 — not in this ADMIN scope row.'
    )
    set_cell_tracked_full(admin_row.cells[2], old_admin_desc, new_admin_desc)

    notes_row = read_all.rows[8]
    old_notes = notes_row.cells[2].text if len(notes_row.cells) > 2 else notes_row.cells[0].text
    extra = (
        ' Rollups visibility requires READ_ALL with non-USER scope. '
        'Class/session reach remains governed by SCHOOL_CLASSES / SCHOOL_SESSIONS.'
    )
    set_cell_tracked_full(
        notes_row.cells[2 if len(notes_row.cells) > 2 else 0],
        old_notes,
        old_notes + extra,
    )

    # --- Table 5: UPDATE ADMIN rows ---
    update_table = tables[5]
    for ri in (8, 9):
        old_text = update_table.rows[ri].cells[2].text
        new_text = old_text.replace('Edit attendance after completed-session deadline.', '').replace('  ', ' ').strip()
        if 'locked session' not in new_text.lower():
            new_text = '+ Edit attendance on locked session (Matrix + Manage Session).'
        else:
            new_text = '+ Edit attendance on locked session (Matrix + Manage Session). ORGANIZATION scope already covers completed-session deadline edits.'
        set_cell_tracked_full(update_table.rows[ri].cells[2], old_text, new_text)

    notes_row = update_table.rows[10]
    old_notes = notes_row.cells[2].text
    red_part = (
        'Enrollment-office N/A locks are governed by SCHOOL_CLASSES UPDATE admin, not SCHOOL_ATTENDANCES scope. '
        'Excuse marking at ORGANIZATION scope follows UPDATE rows (scope-based model); A6 is N/A in admin override form.'
    )
    strike_fragment = (
        'ORGANIZATION/ ADMIN scopes have access to update the attendance and change the enrollment-office N/A locks'
    )
    if strike_fragment in old_notes:
        replace_in_cell_tracked(notes_row.cells[2], strike_fragment, red_part)
    elif 'enrollment-office N/A locks' in old_notes:
        replace_in_cell_tracked(
            notes_row.cells[2],
            'and change the enrollment-office N/A locks\u2014and\u2014completed-session edit policy.',
            red_part,
        )
    else:
        append_cell_tracked(notes_row.cells[2], red_text=red_part)

    # --- Table 6: DELETE rebuild ---
    rebuild_delete_table(doc, tables[6])
    tables = get_tables(doc)

    # --- Tables 8-9: EXPORT / PRINT split ---
    split_export_print_table(tables[8])
    split_export_print_table(tables[9])

    # --- Table 10: Other notes footnote ---
    other_notes = tables[10]
    add_table_row(other_notes, after_index=2)
    foot_row = other_notes.rows[3]
    set_cell_simple(foot_row.cells[0], 'Runtime note', red=True)
    set_cell_simple(
        foot_row.cells[1],
        'Current runtime: matrix link still checks UPDATE via userCanOpenAttendanceMatrix until READ promotion is implemented.',
        red=True,
    )

    # --- Admin overrides section ---
    override_heading = next(
        p for p in doc.paragraphs
        if p.style.name == 'Heading 3' and p.text.strip() == 'Admin and Other Scopes overrides'
    )
    override_para = override_heading._element.getnext()
    while override_para is not None and not override_para.tag.endswith('p'):
        override_para = override_para.getnext()
    if override_para is not None:
        para = Paragraph(override_para, doc)
        build_admin_overrides_table(doc, para)

    # --- Attendance Report intro related docs ---
    report_intro = next(p for p in doc.paragraphs if 'SCHOOL_ATTENDANCE_REPORT' in p.text and p.style.name == 'Normal')
    insert_paragraph_after(
        report_intro,
        'Related: docs/design_docs/admin-access-types-reference-2026-09-05.md for bypass admin vs ADMIN scope column rules.',
        red=True,
    )

    # --- Report tables (locate by content after overrides table insert) ---
    report_read = find_table_containing(doc, 'Student Attendance Report page only')
    if report_read is None:
        raise RuntimeError('Could not find report READ table')
    split_scope_row(
        report_read,
        2,
        'Open Student Attendance Report page only; no attendance field visibility unless READ_ALL is also granted.',
        'Open Student Attendance Report page only; no additional READ extras at this tier (ADMIN scope / Family B).',
        operation='READ',
    )

    report_read_all = find_table_containing(doc, 'page for generating reports')
    if report_read_all is None:
        raise RuntimeError('Could not find report READ_ALL table')
    row = report_read_all.rows[2]
    old_scope = row.cells[1].text
    old_desc = row.cells[2].text
    set_cell_tracked_full(row.cells[1], old_scope, 'ORGANIZATION')
    set_cell_simple(row.cells[2], old_desc)
    add_table_row(report_read_all, after_index=2)
    admin_row = report_read_all.rows[3]
    set_cell_simple(admin_row.cells[0], 'READ_ALL', red=True)
    set_cell_simple(admin_row.cells[1], 'ADMIN', red=True)
    set_cell_simple(admin_row.cells[2], 'No additional READ_ALL extras beyond ORGANIZATION for report page access.', red=True)

    report_export = find_table_containing(doc, 'Export section and can export the reports')
    if report_export is None:
        raise RuntimeError('Could not find report EXPORT table')
    row = report_export.rows[2]
    old_scope = row.cells[1].text
    old_desc = row.cells[2].text
    set_cell_tracked_full(row.cells[1], old_scope, 'ORGANIZATION')
    set_cell_simple(row.cells[2], old_desc)
    add_table_row(report_export, after_index=2)
    admin_row = report_export.rows[3]
    set_cell_simple(admin_row.cells[0], 'EXPORT', red=True)
    set_cell_simple(admin_row.cells[1], 'ADMIN', red=True)
    set_cell_simple(admin_row.cells[2], 'No additional EXPORT extras beyond ORGANIZATION for report export.', red=True)


def verify_docx(path: Path) -> dict:
    doc = Document(str(path))
    tables = get_tables(doc)
    strike_count = 0
    red_count = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.font.strike:
                            strike_count += 1
                        if run.font.color and run.font.color.rgb == RED:
                            red_count += 1
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.strike:
                strike_count += 1
            if run.font.color and run.font.color.rgb == RED:
                red_count += 1

    admin_ref = next((t for t in get_tables(doc) if t.rows[0].cells[0].text.strip() == 'Topic'), None)
    admin_table_rows = len(admin_ref.rows) if admin_ref is not None else 0
    delete_table = next(
        (
            t for t in get_tables(doc)
            if len(t.rows) > 1 and t.rows[1].cells[0].text.strip() == 'DELETE'
        ),
        None,
    )
    delete_cols = len(delete_table.columns) if delete_table is not None else 0

    return {
        'tables': len(tables),
        'admin_table_rows': admin_table_rows,
        'delete_cols': delete_cols,
        'strike_runs': strike_count,
        'red_runs': red_count,
    }


def main() -> int:
    if not DOCX_PATH.exists():
        raise FileNotFoundError(DOCX_PATH)

    doc = Document(str(DOCX_PATH))
    apply_updates(doc)

    temp_path = DOCX_PATH.with_suffix('.docx.tmp')
    doc.save(str(temp_path))
    try:
        temp_path.replace(DOCX_PATH)
        output_path = DOCX_PATH
    except PermissionError:
        fallback = DOCX_PATH.with_name(DOCX_PATH.stem + '-updated.docx')
        temp_path.replace(fallback)
        output_path = fallback
        print(f'Warning: could not overwrite open file. Saved to: {fallback}')

    stats = verify_docx(output_path)
    print(f'Updated: {output_path}')
    print(f"Tables: {stats['tables']}, admin table rows: {stats['admin_table_rows']}, DELETE cols: {stats['delete_cols']}")
    print(f"Strike runs: {stats['strike_runs']}, red runs: {stats['red_runs']}")
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
