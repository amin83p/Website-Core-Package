#!/usr/bin/env python3
"""Generate Attendance Operation and Scope Capabilities Word document.

Uses the Chat operation-scope docx as a style template so layout matches exactly.
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:
    import subprocess

    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx', '-q'])
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / 'docs' / 'chat-operation-scope-capabilities-2026-09-03.docx'
OUTPUT = ROOT / 'docs' / 'attendance-operation-scope-capabilities-2026-09-04.docx'


def clear_document_body(doc: Document) -> None:
    body = doc.element.body
    for child in list(body):
        if child.tag in (qn('w:p'), qn('w:tbl')):
            body.remove(child)


def add_paragraph(doc: Document, text: str, style: str = 'Normal') -> Paragraph:
    return doc.add_paragraph(text, style=style)


def set_cell_text(cell, text: str) -> None:
    cell.text = text


def add_table(doc: Document, rows: list[tuple[str, str, str]], include_header: bool = False) -> Table:
    headers = ('Operation', 'Scope', 'What user can do and see')
    data_rows = rows
    if include_header:
        table_rows = [headers] + list(data_rows)
    else:
        table_rows = list(data_rows)

    table = doc.add_table(rows=len(table_rows), cols=3)
    table.style = 'Table Grid'

    for ri, row in enumerate(table_rows):
        for ci, value in enumerate(row):
            set_cell_text(table.rows[ri].cells[ci], value)
            if include_header and ri == 0:
                for paragraph in table.rows[ri].cells[ci].paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
    doc.add_paragraph('')
    return table


def add_note_row(table: Table, text: str) -> None:
    row = table.add_row()
    for cell in row.cells:
        set_cell_text(cell, text)


def build_document() -> Path:
    if not TEMPLATE.exists():
        raise FileNotFoundError(f'Missing template docx: {TEMPLATE}')

    doc = Document(str(TEMPLATE))
    clear_document_body(doc)

    add_paragraph(doc, 'Attendance Operation and Scope Capabilities', 'Heading 1')
    add_paragraph(doc, 'Access Profile reference | 4 September 2026')
    add_paragraph(
        doc,
        'This matrix describes the effective Attendance behavior after central access approval. '
        'Every action is rechecked against the target class, session, enrollment window, session lock, '
        'completed-session edit policy, and matrix business rules. Client-side flags improve UX but are not authoritative.'
    )

    add_paragraph(doc, 'SCHOOL_ATTENDANCES Section Pages/Modals', 'Heading 2')
    pages = doc.add_table(rows=3, cols=2)
    pages.style = 'Table Grid'
    page_rows = [
        ('Page/Where', 'What user can do and see'),
        (
            '"Attendance Matrix" page',
            'This is the primary attendance workspace at /school/attendances. Users pick an in-scope active class, '
            'load a date window, view student-by-session cells, mark attendance, timing minutes, notes, comments, '
            'files, rollups, Excel export, and change-log queries. Route gate: SCHOOL_ATTENDANCES UPDATE.'
        ),
        (
            'Manage Session — attendance panel',
            'This is the in-page roster attendance area on /school/classes/:classId/sessions/:sessionId. '
            'Today it saves through SCHOOL_SESSIONS UPDATE and session-editor scope. The external Open Attendance Matrix '
            'link uses SCHOOL_ATTENDANCES access through userCanOpenAttendanceMatrix.'
        ),
    ]
    for ri, (left, right) in enumerate(page_rows):
        pages.rows[ri].cells[0].text = left
        pages.rows[ri].cells[1].text = right
        if ri == 0:
            for cell in pages.rows[ri].cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
    doc.add_paragraph('')

    add_paragraph(doc, 'ADMINS And their privileges:', 'Heading 2')
    admins = doc.add_table(rows=8, cols=2)
    admins.style = 'Table Grid'
    admin_rows = [
        ('ADMIN', 'What user can do and see'),
        (
            'BUILT-IN SUPER USERS',
            'These users are not subject to any access limitations or restrictions. Access checks are skipped for them. '
            'The system only verifies whether the active user has super-user privileges and, if so, allows all requested operations.'
        ),
        (
            'ADMIN GLOBAL',
            'These users are not subject to any access limitations or restrictions. Access checks are skipped for them. '
            'The system only verifies whether the active user has global admin privileges and, if so, allows all requested operations.'
        ),
        (
            'ADMIN GLOBAL IN ORGANIZATION',
            'These users are not subject to any access limitations or restrictions in the data scope of selected organization. '
            'Access checks are skipped for them, but they can only manage the data within the selected organization. '
            'The system only verifies whether the active user has global admin privileges and, if so, allows all requested '
            'operations only in the data scope of the selected organization.'
        ),
        (
            'SCOPED ADMIN TO "GENERAL" CATEGORY',
            'These users are not subject to any access limitations or restrictions to access to the sections are classified in '
            '"GENERAL" category. Since "SCHOOL_ATTENDANCES" section is adjusted in "SCHOOL" category, users who have relevant '
            'administrative categories may bypass normal scope limits when section admin is granted. The system verifies admin '
            'privileges and allows all requested operations when admin access to this section applies.'
        ),
        (
            'SCOPED ADMIN TO "GENERAL" CATEGORY IN ORGANIZATION',
            'These users are not subject to any access limitations or restrictions to access to the sections are classified in '
            '"GENERAL" category; only in the data scope of selected organization. Checks are skipped for them in the data scope '
            'of selected organization when section admin is granted.'
        ),
        (
            'ADMIN ACCESS TO THIS SECTION',
            'These users are not subject to any access limitations or restrictions on SCHOOL_ATTENDANCES. They receive attendance '
            'section admin behavior: manage buttons on the matrix, session-lock override, completed-session attendance edit override, '
            'and excuse marking through userCanMarkAttendanceExcused.'
        ),
        (
            'ADMIN ACCESS TO THIS SECTION IN ORGANIZATION',
            'These users are not subject to any access limitations or restrictions on SCHOOL_ATTENDANCES in the data scope of the '
            'selected organization. They receive the same attendance section admin behavior, limited to that organization.'
        ),
    ]
    for ri, (left, right) in enumerate(admin_rows):
        admins.rows[ri].cells[0].text = left
        admins.rows[ri].cells[1].text = right
        if ri == 0:
            for cell in admins.rows[ri].cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
    doc.add_paragraph('')

    add_paragraph(doc, 'Read Access in Access Profile Definitions', 'Heading 2')

    read_rows = [
        ('READ', 'USER', 'No Access.'),
        (
            'READ',
            'OWNER',
            'On the "Attendance Matrix" page:\n'
            'Not wired to attendance routes today. Intended: read-only matrix visibility for owned classes without cell editing.\n'
            'Current runtime: opening the matrix requires SCHOOL_ATTENDANCES UPDATE even for view-only use.'
        ),
        (
            'READ',
            'OWNER',
            'On Manage Session (attendance panel):\n'
            'Not wired through SCHOOL_ATTENDANCES. Manage Session page load uses SCHOOL_SESSIONS READ_ALL today.'
        ),
        (
            'READ',
            'DEPARTMENT/\nDIVISION',
            'On the "Attendance Matrix" page:\n'
            'Not wired to attendance routes today. Intended: read-only matrix for assigned classes (instructor, deliverer, or owner).\n'
            'Current runtime: opening the matrix requires SCHOOL_ATTENDANCES UPDATE.'
        ),
        (
            'READ',
            'DEPARTMENT/\nDIVISION',
            'On Manage Session (attendance panel):\n'
            'Not wired through SCHOOL_ATTENDANCES. Session visibility uses SCHOOL_SESSIONS scope instead.'
        ),
        (
            'READ',
            'ORGANIZATION/\nADMIN',
            'On the "Attendance Matrix" page:\n'
            'Not wired to attendance routes today. Intended: read-only org-wide matrix.\n'
            'Current runtime: opening the matrix requires SCHOOL_ATTENDANCES UPDATE.'
        ),
        (
            'READ',
            'ORGANIZATION/\nADMIN',
            'On Manage Session (attendance panel):\n'
            'Not wired through SCHOOL_ATTENDANCES. Org-wide session visibility uses SCHOOL_SESSIONS scope instead.'
        ),
    ]
    add_table(doc, read_rows, include_header=True)

    read_all_rows = [
        ('READ_ALL', 'USER', 'No Access.'),
        (
            'READ_ALL',
            'OWNER',
            'On the "Attendance Matrix" page:\n'
            'Does not grant page access by itself. When combined with UPDATE, does not enable attendance-admin manage buttons unless '
            'section admin is also granted.'
        ),
        (
            'READ_ALL',
            'OWNER',
            'On Manage Session (attendance panel):\n'
            'No additional attendance-admin behavior through SCHOOL_ATTENDANCES.'
        ),
        (
            'READ_ALL',
            'DEPARTMENT/\nDIVISION',
            'On the "Attendance Matrix" page:\n'
            'When combined with UPDATE and section admin, enables attendance-admin viewer affordances (manage buttons) for assigned classes. '
            'Does not bypass class scope.'
        ),
        (
            'READ_ALL',
            'DEPARTMENT/\nDIVISION',
            'On Manage Session (attendance panel):\n'
            'No in-page attendance-admin affordances through SCHOOL_ATTENDANCES today.'
        ),
        (
            'READ_ALL',
            'ORGANIZATION/\nADMIN',
            'On the "Attendance Matrix" page:\n'
            'When combined with UPDATE and section admin, enables attendance-admin viewer affordances org-wide. Used by '
            'isAttendancesAdminViewerAsync(user, READ_ALL). Does not bypass class scope.'
        ),
        (
            'READ_ALL',
            'ORGANIZATION/\nADMIN',
            'On Manage Session (attendance panel):\n'
            'Attendance section admin can override completed-session attendance edit windows through SCHOOL_ATTENDANCES admin even '
            'when in-page marking still uses SCHOOL_SESSIONS.'
        ),
    ]
    read_all_table = add_table(doc, read_all_rows, include_header=False)
    add_note_row(
        read_all_table,
        'Note: READ_ALL alone does not open the Attendance Matrix. Users still need SCHOOL_ATTENDANCES UPDATE (or attendance section admin) '
        'for page access. READ_ALL admin primarily controls attendance-admin viewer tools and excuse marking eligibility.'
    )

    create_rows = [
        ('CREATE', 'USER', 'No Access.'),
        (
            'CREATE',
            'OWNER/\nDEPARTMENT/\nDIVISION/',
            'On the "Attendance Matrix" page:\n'
            'CREATE is registered in the section catalog but not bound to attendance HTTP routes. First-time roster rows can be created '
            'implicitly when marking a cell; that behavior is gated by UPDATE, not CREATE.'
        ),
        (
            'CREATE',
            'OWNER/\nDEPARTMENT/\nDIVISION/',
            'On Manage Session (attendance panel):\n'
            'Roster rows are created during session save under SCHOOL_SESSIONS UPDATE, not SCHOOL_ATTENDANCES CREATE.'
        ),
        (
            'CREATE',
            'ORGANIZATION/\nADMIN',
            'On the "Attendance Matrix" page:\n'
            'Same as assignment scopes: implicit roster create via UPDATE only.'
        ),
        (
            'CREATE',
            'ORGANIZATION/\nADMIN',
            'On Manage Session (attendance panel):\n'
            'Same as assignment scopes: roster create via SCHOOL_SESSIONS save only.'
        ),
    ]
    create_table = add_table(doc, create_rows, include_header=False)
    add_note_row(
        create_table,
        'Note: A dedicated CREATE route for attendance records is not implemented. Treat CREATE as catalog intent only until promotion '
        'work wires explicit create semantics.'
    )

    update_rows = [
        ('UPDATE', 'USER', 'No Access.'),
        (
            'UPDATE',
            'OWNER',
            'On the "Attendance Matrix" page:\n'
            'Can open /school/attendances and matrix APIs for owned active classes only. Can mark enabled attendance statuses, timing '
            'minutes, and notes. Cannot set excuse flags unless also attendance section admin. Cannot edit locked sessions unless '
            'attendance section admin. Subject to enrollment window, makeup-required status, enrollment-office N/A locks, rolling '
            'capacity, and completed-session edit policy unless attendance admin override applies.'
        ),
        (
            'UPDATE',
            'OWNER',
            'On Manage Session (attendance panel):\n'
            'Current runtime: in-page attendance save uses SCHOOL_SESSIONS UPDATE with session-editor scope, not this SCHOOL_ATTENDANCES row.'
        ),
        (
            'UPDATE',
            'DEPARTMENT/ DIVISION',
            'On the "Attendance Matrix" page:\n'
            'Same as OWNER scope but limited to assigned classes (active instructor, session deliverer, or class owner). Primary instructor '
            'matrix workflow. Profile limits (maxAttempts, maxTimeMinutes, maxVolumeKB) apply when configured.'
        ),
        (
            'UPDATE',
            'DEPARTMENT/ DIVISION',
            'On Manage Session (attendance panel):\n'
            'Current runtime: requires SCHOOL_SESSIONS UPDATE and session-editor scope, which is stricter than matrix class scope.'
        ),
        (
            'UPDATE',
            'ORGANIZATION/ ADMIN',
            'On the "Attendance Matrix" page:\n'
            'Full org matrix for all active classes. With attendance section admin: manage buttons, session-lock override, completed-session '
            'attendance edit override, and excuse marking (userCanMarkAttendanceExcused).'
        ),
        (
            'UPDATE',
            'ORGANIZATION/ ADMIN',
            'On Manage Session (attendance panel):\n'
            'Can open Attendance Matrix link (userCanOpenAttendanceMatrix). In-page attendance save still uses SCHOOL_SESSIONS today. '
            'Attendance admin override for completed-session attendance edits uses SCHOOL_ATTENDANCES UPDATE admin (canOverrideAttendanceEdit).'
        ),
    ]
    update_table = add_table(doc, update_rows, include_header=False)
    add_note_row(
        update_table,
        'Note: Comments and file uploads accept either SCHOOL_ATTENDANCES UPDATE or SCHOOL_SESSIONS UPDATE on shared endpoints. Matrix cell '
        'saves do not re-run session-editor checks; Manage Session saves do.'
    )

    delete_rows = [
        ('DELETE', 'USER', 'No Access.'),
        (
            'DELETE',
            'OWNER',
            'On the "Attendance Matrix" page:\n'
            'DELETE is registered in the section catalog but not bound to attendance routes. Clearing a cell uses UPDATE to empty or N/A '
            'states, not DELETE.'
        ),
        (
            'DELETE',
            'OWNER',
            'On Manage Session (attendance panel):\n'
            'No DELETE route for roster attendance rows.'
        ),
        (
            'DELETE',
            'DEPARTMENT/\nDIVISION',
            'On the "Attendance Matrix" page:\n'
            'No DELETE route. Cell clearing uses UPDATE within assigned class scope.'
        ),
        (
            'DELETE',
            'DEPARTMENT/\nDIVISION',
            'On Manage Session (attendance panel):\n'
            'No DELETE route.'
        ),
        (
            'DELETE',
            'ORGANIZATION/\nADMIN',
            'On the "Attendance Matrix" page:\n'
            'No DELETE route. Org-wide cell clearing uses UPDATE only.'
        ),
        (
            'DELETE',
            'ORGANIZATION/\nADMIN',
            'On Manage Session (attendance panel):\n'
            'No DELETE route.'
        ),
    ]
    add_table(doc, delete_rows, include_header=False)

    delete_all_rows = [
        ('DELETE_ALL', 'USER', 'No Access.'),
        ('DELETE_ALL', 'OWNER', 'No Access'),
        ('DELETE_ALL', 'DEPARTMENT/\nDIVISION', 'No Access.'),
        ('DELETE_ALL', 'ORGANIZATION/\nADMIN', 'No Access.'),
    ]
    delete_all_table = add_table(doc, delete_all_rows, include_header=False)
    add_note_row(
        delete_all_table,
        'Note: DELETE_ALL is registered in the section catalog but no bulk attendance wipe UI or route exists.'
    )

    configure_rows = [
        ('CONFIGURE', 'USER', 'No Access.'),
        (
            'CONFIGURE',
            'OWNER/\nDEPARTMENT/\nDIVISION/',
            'On the "Attendance Matrix" page:\n'
            'No configure UI on operational attendance pages. Matrix consumes org policy through attendanceMatrixPolicyModel without '
            'exposing settings controls.'
        ),
        (
            'CONFIGURE',
            'ORGANIZATION/\nADMIN',
            'On the "Attendance Matrix" page:\n'
            'CONFIGURE is not wired to SCHOOL_ATTENDANCES routes. Org matrix policy catalog, mark appearance, and completed-session edit '
            'policies are maintained under SCHOOL_SETTINGS and related policy models.'
        ),
    ]
    add_table(doc, configure_rows, include_header=False)

    export_rows = [
        ('EXPORT', 'USER', 'No Access.'),
        (
            'EXPORT',
            'OWNER/\nDEPARTMENT/\nDIVISION/',
            'On the "Attendance Matrix" page:\n'
            'EXPORT is registered in the section catalog but Excel export (/api/export.xlsx) is currently gated by UPDATE, not EXPORT.'
        ),
        (
            'EXPORT',
            'ORGANIZATION/\nADMIN',
            'On the "Attendance Matrix" page:\n'
            'Same as assignment scopes today: export requires SCHOOL_ATTENDANCES UPDATE. Intended future split: EXPORT without UPDATE for '
            'read-only exports.'
        ),
    ]
    add_table(doc, export_rows, include_header=False)

    import_rows = [
        ('IMPORT', 'USER', 'No Access.'),
        ('IMPORT', 'OWNER', 'No Access'),
        ('IMPORT', 'DEPARTMENT/\nDIVISION', 'No Access.'),
        (
            'IMPORT',
            'ORGANIZATION/\nADMIN',
            'On the "Attendance Matrix" page:\n'
            'IMPORT is registered in the section catalog but not implemented. No attendance bulk-import route exists.'
        ),
    ]
    add_table(doc, import_rows, include_header=False)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT))
    return OUTPUT


if __name__ == '__main__':
    path = build_document()
    print(f'Generated: {path}')
