#!/usr/bin/env python3
"""Apply revision tracked edits to attendance operation-scope updated docx."""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import RGBColor
    from docx.text.paragraph import Paragraph
except ImportError:
    import subprocess

    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx', '-q'])
    from docx import Document
    from docx.shared import RGBColor
    from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
HELPER_PATH = ROOT / 'scripts' / 'school' / 'update-attendance-operation-scope-docx.py'
_spec = importlib.util.spec_from_file_location('attendance_docx_helpers', HELPER_PATH)
_helpers = importlib.util.module_from_spec(_spec)
assert _spec.loader is not None
_spec.loader.exec_module(_helpers)

get_tables = _helpers.get_tables
find_table_containing = _helpers.find_table_containing
set_cell_tracked_full = _helpers.set_cell_tracked_full
set_cell_simple = _helpers.set_cell_simple
append_cell_tracked = _helpers.append_cell_tracked
replace_in_cell_tracked = _helpers.replace_in_cell_tracked

DOCX_PATH = ROOT / 'docs' / 'attendance-operation-scope-capabilities-2026-09-04-updated.docx'

READ_ALL_ADMIN_NEW = (
    'On the "Attendance Matrix" page:\n'
    'No additional READ_ALL extras beyond ORGANIZATION.\n'
    'Locked sessions remain visible to all users who can open the matrix; only cell updates on locked '
    'sessions are blocked (unless UPDATE + ADMIN scope or bypass override A1).\n'
    'Admin viewer tools and full field visibility are bypass-only — see overrides A4 and A12.'
)

ROLLUPS_SCOPE_OLD_FRAGMENTS = [
    'Rollups visibility requires READ_ALL with non-USER and OWNER scopes.',
    'Rollups visibility requires READ_ALL with non-USER\u00a0scope.',
    'Rollup columns/values are returned and displayed only when the user also has\u00a0READ_ALL\u00a0with a non-USER\u00a0scope.',
]

ROLLUPS_SCOPE_NEW = (
    'Rollups visibility requires READ_ALL with DEPARTMENT/DIVISION, ORGANIZATION, or ADMIN scope. '
    'USER and OWNER: No Access to rollups.'
)

ROLLUPS_API_NEW = (
    'POST /api/rollups: Requires SCHOOL_ATTENDANCES READ to call the API. '
    'Rollup columns/values are displayed only when the user also has READ_ALL with '
    'DEPARTMENT/DIVISION, ORGANIZATION, or ADMIN scope. USER and OWNER: No Access to rollups.'
)

UPDATE_NOTES_NEW = (
    'Notes:\n'
    '- Comments on shared endpoints accept SCHOOL_ATTENDANCES UPDATE or SCHOOL_SESSIONS UPDATE. '
    'Matrix cell saves do not re-run session-editor checks; Manage Session saves do.\n'
    '- File uploads on shared endpoints require SCHOOL_ATTENDANCES UPLOAD (preferred) or '
    'SCHOOL_SESSIONS UPLOAD when attendance UPLOAD is not granted.\n'
    '- Locked sessions remain visible in the matrix for all scopes that can open the attendance context; '
    'lock blocks UPDATE only (not READ or READ_ALL visibility).'
)

LOCKED_SESSION_NOTE = (
    'Locked sessions remain visible in the matrix for all scopes that can open the attendance context; '
    'lock blocks UPDATE only (not READ or READ_ALL visibility).'
)


def find_read_all_matrix_table(doc: Document):
    for table in get_tables(doc):
        if len(table.rows) < 8:
            continue
        if table.rows[1].cells[0].text.strip() != 'READ_ALL':
            continue
        if table.rows[7].cells[1].text.strip() == 'ADMIN':
            if 'Attendance Matrix' in table.rows[2].cells[2].text or 'Attendance Matrix' in table.rows[7].cells[2].text:
                return table
    return None


def find_print_matrix_table(doc: Document):
    for table in get_tables(doc):
        if len(table.rows) < 3:
            continue
        if table.rows[1].cells[0].text.strip() != 'PRINT':
            continue
        if 'Attendance Matrix' in table.rows[2].cells[2].text:
            return table
    return None


def find_update_table(doc: Document):
    for table in get_tables(doc):
        if len(table.rows) < 9:
            continue
        if table.rows[1].cells[0].text.strip() != 'UPDATE':
            continue
        if 'Notes' in table.rows[-1].cells[0].text or 'Notes' in table.rows[-1].cells[2].text:
            return table
    return None


def fix_rollups_text(cell) -> None:
    text = cell.text
    for fragment in ROLLUPS_SCOPE_OLD_FRAGMENTS:
        if fragment in text:
            replace_in_cell_tracked(cell, fragment, ROLLUPS_SCOPE_NEW)
            return
    if 'non-USER' in text and 'OWNER' not in text.split('non-USER')[1][:30]:
        replace_in_cell_tracked(cell, 'with a non-USER scope.', ROLLUPS_SCOPE_NEW)
        return
    if 'USER and OWNER' not in text:
        append_cell_tracked(cell, red_text=ROLLUPS_SCOPE_NEW)


def insert_upload_table_after_print(doc: Document) -> None:
    if find_table_containing(doc, 'No additional UPLOAD extras beyond ORGANIZATION'):
        return

    print_table = find_print_matrix_table(doc)
    if print_table is None:
        raise RuntimeError('PRINT table not found')

    table = doc.add_table(rows=4, cols=3)
    table.style = 'Table Grid'
    tbl_element = table._element
    tbl_element.getparent().remove(tbl_element)
    print_table._element.addnext(tbl_element)
    headers = ['Operation', 'Scope', 'What user can do and see']
    for ci, header in enumerate(headers):
        set_cell_simple(table.rows[0].cells[ci], header)
        for run in table.rows[0].cells[ci].paragraphs[0].runs:
            run.bold = True

    rows = [
        ('UPLOAD', 'USER / OWNER', 'No Access.'),
        (
            'UPLOAD',
            'DEPARTMENT / DIVISION / ORGANIZATION',
            'On the "Attendance Matrix" page: + Can upload files attached to attendance details.\n'
            'On Manage Session (attendance panel): + Can upload files attached to roster attendance details.',
        ),
        ('UPLOAD', 'ADMIN', 'No additional UPLOAD extras beyond ORGANIZATION.'),
    ]
    for ri, (op, scope, desc) in enumerate(rows, start=1):
        set_cell_simple(table.rows[ri].cells[0], op, red=True)
        set_cell_simple(table.rows[ri].cells[1], scope, red=True)
        set_cell_simple(table.rows[ri].cells[2], desc, red=True)

    doc.add_paragraph('')


def apply_revision(doc: Document) -> None:
    read_all = find_read_all_matrix_table(doc)
    if read_all is None:
        raise RuntimeError('READ_ALL matrix table not found')

    admin_row = read_all.rows[7]
    set_cell_tracked_full(admin_row.cells[2], admin_row.cells[2].text, READ_ALL_ADMIN_NEW)

    notes_row = read_all.rows[8]
    notes_cell = notes_row.cells[2] if len(notes_row.cells) > 2 else notes_row.cells[0]
    old_notes = notes_cell.text
    fix_rollups_text(notes_cell)
    if LOCKED_SESSION_NOTE not in notes_cell.text:
        append_cell_tracked(notes_cell, red_text=LOCKED_SESSION_NOTE)

    update_table = find_update_table(doc)
    if update_table is None:
        raise RuntimeError('UPDATE table not found')

    admin_update = update_table.rows[8]
    if admin_update.cells[1].text.strip() == 'ADMIN':
        old_admin = admin_update.cells[2].text
        new_admin = old_admin.replace('..', '.').replace(
            'On the "Attendance Matrix" page and',
            'On the "Attendance Matrix" page and',
        )
        if new_admin != old_admin or '..' in old_admin:
            set_cell_tracked_full(admin_update.cells[2], old_admin, new_admin.replace('..', '.'))

    notes_update = update_table.rows[-1]
    notes_cell = notes_update.cells[2] if len(notes_update.cells) > 2 else notes_update.cells[0]
    set_cell_tracked_full(notes_cell, notes_cell.text, UPDATE_NOTES_NEW)

    other_notes = find_table_containing(doc, 'Open matrix link from Manage Session')
    if other_notes is not None and len(other_notes.rows) > 2:
        rollups_cell = other_notes.rows[2].cells[1]
        old_rollups = rollups_cell.text
        set_cell_tracked_full(rollups_cell, old_rollups, ROLLUPS_API_NEW)

    insert_upload_table_after_print(doc)


def main() -> int:
    if not DOCX_PATH.exists():
        raise FileNotFoundError(DOCX_PATH)

    doc = Document(str(DOCX_PATH))
    apply_revision(doc)

    temp_path = DOCX_PATH.with_suffix('.docx.tmp')
    doc.save(str(temp_path))
    try:
        temp_path.replace(DOCX_PATH)
        output = DOCX_PATH
    except PermissionError:
        output = DOCX_PATH.with_name(DOCX_PATH.stem + '-rev2.docx')
        temp_path.replace(output)

    doc2 = Document(str(output))
    tables = get_tables(doc2)
    read_all = find_read_all_matrix_table(doc2)
    upload = find_table_containing(doc2, 'No additional UPLOAD extras beyond ORGANIZATION')
    print(f'Updated: {output}')
    print(f'Tables: {len(tables)}, UPLOAD table: {"yes" if upload else "no"}')
    if read_all:
        print('READ_ALL ADMIN:', read_all.rows[7].cells[2].text[:120].replace('\n', ' '))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
