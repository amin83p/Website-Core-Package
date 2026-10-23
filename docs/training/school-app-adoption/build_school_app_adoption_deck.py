#!/usr/bin/env python3
"""
Build school app adoption training PowerPoint and one-page teacher handout PDF.
Replace files in screenshots/ with real captures (same filenames) and re-run.
"""
from __future__ import annotations

import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent
SCREENSHOTS = ROOT / "screenshots"
MANIFEST = ROOT / "screenshot-manifest.json"
OUT_PPTX = ROOT / "School-App-Adoption-Training.pptx"
OUT_PDF = ROOT / "School-App-Adoption-Handout.pdf"

ACCENT = RGBColor(0x0D, 0x6E, 0xFD)
ACCENT_HEX = "#0d6efd"
DARK = RGBColor(0x21, 0x25, 0x29)
MUTED = RGBColor(0x6C, 0x75, 0x7D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF3, 0xF8, 0xFF)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def load_manifest() -> dict:
    with MANIFEST.open(encoding="utf-8") as f:
        return json.load(f)


def try_font(size: int, bold: bool = False):
    names = ["segoeui.ttf", "Segoe UI.ttf", "arial.ttf", "Arial.ttf"]
    if bold:
        names = ["segoeuib.ttf", "Segoe UI Bold.ttf", "arialbd.ttf"] + names
    for name in names:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def ensure_placeholder_screenshots(manifest: dict) -> None:
    SCREENSHOTS.mkdir(parents=True, exist_ok=True)
    title_font = try_font(36, bold=True)
    hint_font = try_font(22)
    num_font = try_font(28, bold=True)
    sub_font = try_font(18)

    for idx, shot in enumerate(manifest["shots"], start=1):
        path = SCREENSHOTS / shot["file"]
        if path.exists() and path.stat().st_size > 50_000:
            continue
        w, h = 1600, 900
        img = Image.new("RGB", (w, h), "#f8fbff")
        draw = ImageDraw.Draw(img)
        draw.rounded_rectangle((40, 40, w - 40, h - 40), radius=24, outline=ACCENT_HEX, width=4)
        draw.rounded_rectangle((60, 60, w - 60, h - 60), radius=16, fill="#ffffff", outline="#dce9ff", width=2)

        draw.text((80, 80), f"Screenshot {idx:02d}", fill=ACCENT_HEX, font=num_font)
        draw.text((80, 140), shot["title"], fill="#212529", font=title_font)
        draw.text((80, 210), shot["hint"], fill="#6c757d", font=hint_font)

        wrap_y = 280
        lines = [
            "Replace this file with your capture.",
            f"Filename: {shot['file']}",
            "Blur student names if shown.",
        ]
        if shot.get("needsUserCapture"):
            lines.append("Needs real school data — see README.")
        for line in lines:
            draw.text((80, wrap_y), line, fill="#495057", font=sub_font)
            wrap_y += 32

        draw.text((80, h - 100), f"Slide {shot['slide']}", fill=ACCENT_HEX, font=sub_font)
        img.save(path, "PNG", optimize=True)


def set_slide_bg(slide, prs: Presentation, fill: RGBColor = WHITE) -> None:
    fill_el = slide.background.fill
    fill_el.solid()
    fill_el.fore_color.rgb = fill


def add_footer_bar(slide, prs: Presentation, text: str = "School app adoption") -> None:
    shape = slide.shapes.add_shape(1, Inches(0), SLIDE_H - Inches(0.35), prs.slide_width, Inches(0.35))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()
    tf = shape.text_frame
    tf.text = text
    p = tf.paragraphs[0]
    p.font.size = Pt(11)
    p.font.color.rgb = WHITE
    p.font.name = "Segoe UI"
    p.alignment = PP_ALIGN.CENTER


def add_title_block(slide, title: str, subtitle: str | None = None, top: float = 0.45) -> None:
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12.1), Inches(1.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = DARK
    p.font.name = "Segoe UI"
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(18)
        p2.font.color.rgb = MUTED
        p2.font.name = "Segoe UI"
        p2.space_before = Pt(6)


def add_bullets(
    slide,
    items: list[str],
    left: float = 0.65,
    top: float = 1.55,
    width: float = 5.8,
    size: int = 20,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.2))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = DARK
        p.font.name = "Segoe UI"
        p.space_after = Pt(10)
        if item.strip():
            p.bullet = True


def add_image(slide, image_path: Path, left: float, top: float, width: float, height: float) -> None:
    if image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(left), Inches(top), width=Inches(width), height=Inches(height))


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def shots_map(manifest: dict) -> dict[int, list[Path]]:
    out: dict[int, list[Path]] = {}
    for shot in manifest["shots"]:
        out.setdefault(shot["slide"], []).append(SCREENSHOTS / shot["file"])
    return out


def slide_text_only(prs, blank, title: str, bullets: list[str], notes: str, footer: str | None = None) -> None:
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, prs)
    add_title_block(s, title)
    add_bullets(s, bullets, width=11.5, top=1.65)
    add_footer_bar(s, prs, footer or "School app adoption")
    add_notes(s, notes)


def slide_with_shots(
    prs,
    blank,
    title: str,
    bullets: list[str],
    notes: str,
    by_slide: dict[int, list[Path]],
    slide_num: int,
    subtitle: str | None = None,
    bullet_width: float = 5.6,
    bullet_size: int = 19,
    one_image: bool = True,
) -> None:
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, prs)
    add_title_block(s, title, subtitle)
    add_bullets(s, bullets, width=bullet_width, size=bullet_size)
    imgs = by_slide.get(slide_num, [])
    if one_image and imgs:
        add_image(s, imgs[0], 6.3, 1.35, 6.4, 5.15)
    add_footer_bar(s, prs)
    add_notes(s, notes)


def slide_two_images_bottom(
    prs,
    blank,
    title: str,
    bullets: list[str],
    notes: str,
    by_slide: dict[int, list[Path]],
    slide_num: int,
) -> None:
    s = prs.slides.add_slide(blank)
    set_slide_bg(s, prs)
    add_title_block(s, title)
    add_bullets(s, bullets, width=5.6, size=18)
    imgs = by_slide.get(slide_num, [])
    if len(imgs) >= 1:
        add_image(s, imgs[0], 6.25, 1.35, 3.1, 2.5)
    if len(imgs) >= 2:
        add_image(s, imgs[1], 9.55, 1.35, 3.1, 2.5)
    add_footer_bar(s, prs)
    add_notes(s, notes)


def add_audience_columns(slide, prs: Presentation) -> None:
    groups = [
        ("Teachers", ["See today's sessions", "One Manage Session workspace", "Clear completion status"]),
        ("Staff", ["Enrollment & records", "Less chasing missing data", "Shared student picture"]),
        ("Managers", ["Visibility & reports", "Policies & access windows", "Oversight without micromanaging"]),
    ]
    x0 = 0.7
    for i, (heading, lines) in enumerate(groups):
        left = x0 + i * 4.15
        box = slide.shapes.add_shape(1, Inches(left), Inches(2.0), Inches(3.85), Inches(4.35))
        box.fill.solid()
        box.fill.fore_color.rgb = LIGHT_BG
        box.line.color.rgb = ACCENT
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = heading
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = ACCENT
        p.font.name = "Segoe UI"
        for line in lines:
            p2 = tf.add_paragraph()
            p2.text = line
            p2.font.size = Pt(16)
            p2.font.color.rgb = DARK
            p2.font.name = "Segoe UI"
            p2.space_before = Pt(8)
            p2.bullet = True


def build_presentation(manifest: dict) -> Presentation:
    school = manifest.get("schoolName", "Your School Name")
    completed = manifest.get("completedStatusLabel", "Completed")
    support = manifest.get("supportContact", "Contact your school office or IT support.")
    by_slide = shots_map(manifest)

    prs = Presentation()
    prs.slide_width = int(SLIDE_W)
    prs.slide_height = int(SLIDE_H)
    blank = prs.slide_layouts[6]

    # 1 Title
    s1 = prs.slides.add_slide(blank)
    set_slide_bg(s1, prs, LIGHT_BG)
    accent = s1.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.25), SLIDE_H)
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT
    accent.line.fill.background()
    tbox = s1.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(11.5), Inches(3.0))
    tf = tbox.text_frame
    p = tf.paragraphs[0]
    p.text = "Why we use the school app"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = DARK
    p.font.name = "Segoe UI"
    p2 = tf.add_paragraph()
    p2.text = school
    p2.font.size = Pt(24)
    p2.font.color.rgb = ACCENT
    p2.font.name = "Segoe UI"
    p2.space_before = Pt(14)
    p3 = tf.add_paragraph()
    p3.text = "Migration · Teachers · Staff · Leadership · Daily routine"
    p3.font.size = Pt(16)
    p3.font.color.rgb = MUTED
    p3.font.name = "Segoe UI"
    p3.space_before = Pt(24)
    add_footer_bar(s1, prs, "Training deck — replace screenshot placeholders with your captures")
    add_notes(s1, "Goals: why we adopted the app, who it helps, and what teachers do every teaching day.")

    slide_text_only(
        prs,
        blank,
        "One place for the school day",
        [
            "Single source of truth for schedules, sessions, and records",
            "Less duplicate spreadsheets, email threads, and paper handoffs",
            "Everyone works from the same dates, classes, and session statuses",
        ],
        "Emphasize consistency: when data lives in one system, fewer surprises for families and staff.",
    )

    slide_text_only(
        prs,
        blank,
        "Accountability and access",
        [
            "Role-based permissions — teachers see their work; offices see what they need",
            "Changes are tied to accounts, not anonymous files",
            "Audit-friendly history for attendance and session completion",
        ],
        "You are not being watched for micromanagement — the system protects everyone when questions come up later.",
    )

    slide_text_only(
        prs,
        blank,
        "Better outcomes",
        [
            "Continuity for students when staff or teachers change",
            "Compliance with school policies and reporting requirements",
            "Cleaner handoff to payroll and timesheet processes where linked",
        ],
        "Keep this high level; policy details live with your admin team.",
    )

    slide_text_only(
        prs,
        blank,
        'What "good" looks like',
        [
            "Every taught session is opened in Manage Session",
            "Attendance, notes, and assigned reports are filled",
            f"Session status moves to {completed} and Save session is used",
        ],
        "Good is complete sessions, not perfect prose in every note field.",
    )

    slide_text_only(
        prs,
        blank,
        "Why we moved to this app",
        [
            "Scheduling, session work, and attendance live in one platform",
            "Security and hosting aligned with school IT standards",
            "One codebase we can improve — fixes and features benefit everyone",
        ],
        "Migration is about sustainability and reliability, not change for its own sake.",
    )

    slide_text_only(
        prs,
        blank,
        "What changes for you",
        [
            "Your teaching in the classroom stays the same",
            "You record the day in the app instead of scattered tools",
            "Daily habit: School dashboard → Schedule Overview → each session",
        ],
        "Acknowledge learning curve; the daily path is short and repeatable.",
    )

    # 8 Three groups
    s8 = prs.slides.add_slide(blank)
    set_slide_bg(s8, prs)
    add_title_block(s8, "Three groups, one system")
    add_audience_columns(s8, prs)
    add_footer_bar(s8, prs)
    add_notes(s8, "Same data, different lenses — teachers enter, staff operate, managers oversee.")

    slide_text_only(
        prs,
        blank,
        "How the app helps teachers",
        [
            "Schedule Overview shows what you teach and what's already done",
            "Manage Session is one workspace per class meeting",
            f"{completed} status makes 'caught up' visible to you and the office",
        ],
        "Point to the short daily deck in docs/training/teacher-daily-app for a refresher.",
    )

    slide_text_only(
        prs,
        blank,
        "How the app helps staff",
        [
            "Enrollment, finance, library, and communications pull from shared data",
            "Fewer calls to chase attendance sheets or session paperwork",
            "Student cases and reports arrive in context of the right session",
        ],
        "Staff slides are overview only — their training is role-specific.",
    )

    slide_text_only(
        prs,
        blank,
        "How the app helps leadership",
        [
            "Dashboards and reports without opening every teacher's files",
            "Session access and completion policies enforced in the app",
            "Support teachers with clear expectations, not cell-by-cell policing",
        ],
        "Leadership cares about completion rates and audit readiness.",
    )

    # 12 Divider
    s12 = prs.slides.add_slide(blank)
    set_slide_bg(s12, prs, LIGHT_BG)
    add_title_block(s12, "Your daily routine", "Every teaching day", top=2.4)
    sub = s12.shapes.add_textbox(Inches(0.65), Inches(3.5), Inches(12), Inches(1.5))
    sub.text_frame.text = "Open Schedule Overview → each session → Manage Session → save → verify"
    sub.text_frame.paragraphs[0].font.size = Pt(22)
    sub.text_frame.paragraphs[0].font.color.rgb = MUTED
    sub.text_frame.paragraphs[0].font.name = "Segoe UI"
    add_footer_bar(s12, prs)
    add_notes(s12, "Transition to hands-on walkthrough with screenshots.")

    slide_with_shots(
        prs,
        blank,
        "Start on the dashboard",
        [
            "Log in and open the School section",
            "Find My Schedule Overview",
            "Click Open My Schedule",
        ],
        "This is the front door for teachers.",
        by_slide,
        13,
        subtitle="School dashboard",
    )

    # 14 two images
    s14 = prs.slides.add_slide(blank)
    set_slide_bg(s14, prs)
    add_title_block(s14, "Open Schedule Overview every day", "What am I teaching, when, and what's done?")
    add_bullets(
        s14,
        [
            "Page title: Schedule Overview",
            "Set Anchor Date — usually today",
            "Choose period (Daily, Weekly, …) → Load",
            "Scan Session Details for status colors",
        ],
        width=5.5,
        size=18,
    )
    imgs14 = by_slide.get(14, [])
    if imgs14:
        add_image(s14, imgs14[0], 6.3, 1.35, 3.2, 5.15)
    if len(imgs14) > 1:
        add_image(s14, imgs14[1], 9.65, 1.35, 3.05, 5.15)
    add_footer_bar(s14, prs)
    add_notes(s14, "Weekly + today's anchor for planning; Daily on teaching days.")

    slide_with_shots(
        prs,
        blank,
        "Pick your session",
        [
            "Each row = one class session",
            "Confirm Role shows Teacher for sessions you lead",
            "Click the row to open Manage Session",
        ],
        "One session at a time — repeat for every class today.",
        by_slide,
        15,
        subtitle="Session Details",
    )

    # 16 two images
    slide_two_images_bottom(
        prs,
        blank,
        "Review session details",
        [
            "Check date, time, room, and roster context",
            "Note current Session status before you edit",
            "Use the left menu to move between work areas",
            "Save session often while you work",
        ],
        "Tour the layout before diving into attendance.",
        by_slide,
        16,
    )

    slide_with_shots(
        prs,
        blank,
        "Mark attendance",
        [
            "Open the Attendance panel",
            "Mark Present, Late, Absent, and other codes your school uses",
            "Use Open attendance matrix when your school links matrix edits",
        ],
        "Blur student names in training screenshots.",
        by_slide,
        17,
    )

    slide_with_shots(
        prs,
        blank,
        "Session notes",
        [
            'Open Session & notes',
            "Record what happened instructionally",
            "Enough detail for the next teacher or admin follow-up",
        ],
        "Skip only if your school explicitly allows empty notes — default is fill them.",
        by_slide,
        18,
    )

    slide_with_shots(
        prs,
        blank,
        "Assigned reports",
        [
            "Open Reports in the sidebar",
            "Complete every report assigned to you for this session",
            "Save session after filling reports",
        ],
        "If no assignments show, your school may not use session reports yet.",
        by_slide,
        19,
    )

    slide_with_shots(
        prs,
        blank,
        "Book coverage",
        [
            "Open Book coverage",
            "Log books or materials used in this session",
            "Required where your program tracks library or curriculum resources",
        ],
        "Skip when no book was used — only when truly N/A.",
        by_slide,
        20,
    )

    slide_two_images_bottom(
        prs,
        blank,
        "Student and class records",
        [
            "Student Cases — per-student follow-ups tied to the session",
            "Class Conduct — class-wide conduct when your school uses it",
            "Report incidents and cases when policy requires it",
        ],
        "Not every session needs cases — use when something must be documented.",
        by_slide,
        21,
    )

    slide_with_shots(
        prs,
        blank,
        "Activities and scores",
        [
            "Open Grade book when enabled for your classes",
            "Record scored and unscored activities as your program requires",
            "Aligns student progress with the session date",
        ],
        "If Grade book is disabled or 'coming soon', skip and note for your admin.",
        by_slide,
        22,
    )

    slide_with_shots(
        prs,
        blank,
        f"Finish the session — {completed}",
        [
            "Open Session status at the top of Manage Session",
            f"Choose {completed}",
            "Look for the session completed indicator when shown",
        ],
        "Session status is separate from attendance — both may be required.",
        by_slide,
        23,
    )

    # 24 two images
    s24 = prs.slides.add_slide(blank)
    set_slide_bg(s24, prs)
    add_title_block(s24, "Save and verify")
    add_bullets(
        s24,
        [
            "Click Save session",
            "Return to Schedule Overview → Load or refresh",
            f"Confirm the row shows {completed}",
        ],
        width=5.6,
    )
    imgs24 = by_slide.get(24, [])
    if len(imgs24) >= 1:
        add_image(s24, imgs24[0], 6.25, 1.35, 3.1, 2.5)
    if len(imgs24) >= 2:
        add_image(s24, imgs24[1], 9.55, 1.35, 3.1, 2.5)
    add_footer_bar(s24, prs)
    add_notes(s24, "Verification is how you know the week is caught up.")

    # 25 Recap
    s25 = prs.slides.add_slide(blank)
    set_slide_bg(s25, prs, LIGHT_BG)
    add_title_block(s25, "Daily checklist + Q&A", top=0.55)
    steps = [
        "Open My Schedule from the School dashboard",
        "Schedule Overview: Anchor Date + period → Load",
        "Open each session → Manage Session",
        "Attendance · Notes · Reports · Book coverage · Cases · Grade book (as shown)",
        f"Session status = {completed} → Save session → verify on schedule",
    ]
    y = 1.75
    for i, step in enumerate(steps, start=1):
        circ = s25.shapes.add_shape(9, Inches(0.75), Inches(y), Inches(0.55), Inches(0.55))
        circ.fill.solid()
        circ.fill.fore_color.rgb = ACCENT
        circ.line.fill.background()
        tf = circ.text_frame
        tf.text = str(i)
        tf.paragraphs[0].font.size = Pt(18)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = WHITE
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tb = s25.shapes.add_textbox(Inches(1.45), Inches(y - 0.05), Inches(10.5), Inches(0.65))
        tb.text_frame.text = step
        tb.text_frame.paragraphs[0].font.size = Pt(20)
        tb.text_frame.paragraphs[0].font.color.rgb = DARK
        tb.text_frame.paragraphs[0].font.name = "Segoe UI"
        y += 0.88
    help_box = s25.shapes.add_shape(1, Inches(0.65), Inches(6.35), Inches(11.9), Inches(0.75))
    help_box.fill.solid()
    help_box.fill.fore_color.rgb = WHITE
    help_box.line.color.rgb = ACCENT
    help_box.text_frame.text = support
    help_box.text_frame.paragraphs[0].font.size = Pt(16)
    help_box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    help_box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    add_footer_bar(s25, prs, "Handout PDF matches this checklist")
    add_notes(s25, "Leave time for questions; offer 5 minutes hands-on if in person.")

    return prs


def build_handout_pdf(manifest: dict) -> None:
    school = manifest.get("schoolName", "Your School Name")
    completed = manifest.get("completedStatusLabel", "Completed")
    support = manifest.get("supportContact", "Contact your school office or IT support.")
    c = canvas.Canvas(str(OUT_PDF), pagesize=letter)
    w, h = letter
    margin = 0.75 * inch

    c.setFillColor(colors.HexColor(ACCENT_HEX))
    c.rect(0, h - 0.45 * inch, w, 0.45 * inch, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 11)
    c.drawCentredString(w / 2, h - 0.32 * inch, f"{school} — teacher daily checklist")

    c.setFillColor(colors.HexColor("#212529"))
    c.setFont("Helvetica-Bold", 20)
    c.drawString(margin, h - 1.1 * inch, "School app — daily teacher routine")

    c.setFont("Helvetica", 11)
    c.setFillColor(colors.HexColor("#6c757d"))
    c.drawString(margin, h - 1.38 * inch, "Schedule Overview → Manage Session → Save → verify")

    steps = [
        "School dashboard → My Schedule Overview → Open My Schedule",
        "Schedule Overview: set Anchor Date, choose period, click Load",
        "Click each session you teach → Manage Session",
        "Attendance — mark all students",
        "Session & notes — record what happened",
        "Reports — complete assignments for this session",
        "Book coverage — log books/materials used",
        "Student Cases and Class Conduct — when required",
        "Grade book — scored/unscored activities as required",
        f"Session status = {completed} → Save session",
        f"Back on Schedule Overview — row shows {completed}",
    ]
    y = h - 1.85 * inch
    c.setFillColor(colors.HexColor("#212529"))
    for i, step in enumerate(steps, start=1):
        c.setFillColor(colors.HexColor(ACCENT_HEX))
        c.circle(margin + 8, y + 4, 10, fill=1, stroke=0)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 9)
        c.drawCentredString(margin + 8, y, str(i))
        c.setFillColor(colors.HexColor("#212529"))
        c.setFont("Helvetica", 11)
        c.drawString(margin + 26, y - 2, step)
        y -= 0.48 * inch
        if y < 1.2 * inch:
            c.showPage()
            y = h - 1 * inch

    c.setFillColor(colors.HexColor("#6c757d"))
    c.setFont("Helvetica-Oblique", 9)
    c.drawString(margin, 0.85 * inch, support)

    c.showPage()
    c.save()


def main() -> None:
    manifest = load_manifest()
    ensure_placeholder_screenshots(manifest)
    prs = build_presentation(manifest)
    prs.save(str(OUT_PPTX))
    build_handout_pdf(manifest)
    print(f"Wrote {OUT_PPTX}")
    print(f"Wrote {OUT_PDF}")
    print(f"Screenshots in {SCREENSHOTS}")


if __name__ == "__main__":
    main()
